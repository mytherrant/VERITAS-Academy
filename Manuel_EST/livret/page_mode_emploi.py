# -*- coding: utf-8 -*-
"""
Ajoute en tête de chaque LIVRET D'ACTIVITÉS la page « Comment utiliser ce livret ».
S'exécute APRÈS build_livret.py.

Usage :  python page_mode_emploi.py <2nde|1ere|Tle>
"""
import sys, os, copy
from docx import Document
from docx.text.paragraph import Paragraph

MARQUE = "Comment utiliser ce livret"

PAGE = [
    ('Rubrique',      "🧭 Comment utiliser ce livret"),
    ('Corps',         "Ce livret ne te donne pas les réponses : il te donne la place de les trouver. "
                      "Chaque règle que tu liras ici, c'est toi qui l'écriras. Ce que tu écris de ta "
                      "main, tu le retiens ; ce que tu recopies d'un corrigé, tu l'oublies avant l'examen."),
    ('Retient_Titre', "🗓 Les quatre temps d'une leçon"),
    ('Liste',         "⏱ AVANT LE COURS — chez moi, seul(e), au crayon. Je lis le corpus et je réponds "
                      "aux questions du « Je manipule ». Je me trompe : c'est normal, c'est ainsi qu'on apprend."),
    ('Liste',         "👥 EN CLASSE — à deux. Je compare mes réponses avec celles de mon voisin. "
                      "Chaque réponse s'appuie sur un mot du texte : je le recopie."),
    ('Liste',         "🗣 EN CLASSE — avec le professeur. Nous formulons ensemble la règle. "
                      "Je remplis alors le cadre « Je retiens », au stylo."),
    ('Liste',         "🏠 APRÈS LE COURS — je fais les deux exercices. Ils ne répètent pas la leçon : "
                      "ils la mettent à l'épreuve sur un texte nouveau."),
    ('Retient_Titre', "🖊 Les signes du livret"),
    ('Liste',         "🎯 Mon objectif — les verbes sont en pointillés : je les compléterai en fin de "
                      "leçon, quand je saurai ce que je sais faire."),
    ('Liste',         "📌 Corpus — le texte d'auteur. Il est donné en entier : on ne résume jamais "
                      "ce qu'on n'a pas lu."),
    ('Liste',         "🧠 Je retiens — un cadre à remplir, jamais une leçon à lire."),
    ('Liste',         "✍️ Je m'exerce — deux exercices ouverts. Les corrigés sont en ligne sur "
                      "veritas-school.com, à consulter après avoir cherché."),
    ('Astuce',        "💡 Côté enseignant — Ce livret est conçu pour la classe inversée. Les cadres "
                      "« Je retiens » se remplissent en classe, après la phase de manipulation : c'est "
                      "la trace écrite de la leçon, pas son point de départ. Le cahier corrigé "
                      "(version enseignant) contient les mêmes leçons, rédigées."),
]


def ajouter(src, dst=None):
    """Insère la page après le sous-titre de couverture. Idempotent."""
    d = Document(src)
    if any(MARQUE in p.text for p in d.paragraphs[:40]):
        return 0                      # déjà posée : on ne double pas
    ancre = d.paragraphs[1]           # juste après le sous-titre de couverture
    for style, txt in PAGE:
        new = copy.deepcopy(ancre._element)
        ancre._element.addnext(new)
        np = Paragraph(new, ancre._parent)
        for r in np.runs[1:]:
            r._element.getparent().remove(r._element)
        try:
            np.style = style
        except KeyError:
            pass
        if np.runs:
            np.runs[0].text = txt
        else:
            np.add_run(txt)
        ancre = np
    d.save(dst or src)
    return len(PAGE)


if __name__ == '__main__':
    niv = sys.argv[1] if len(sys.argv) > 1 else '2nde'
    f = os.path.join(os.path.dirname(__file__), '..', 'cahiers',
                     'Cahier_%s_ST_LIVRET_ACTIVITES.docx' % niv)
    n = ajouter(f)
    print('%s : %s' % (os.path.basename(f),
                       'page mode d\'emploi ajoutée (%d §)' % n if n
                       else 'page déjà présente, inchangée'))
