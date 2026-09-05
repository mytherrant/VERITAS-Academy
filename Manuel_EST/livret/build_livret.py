# -*- coding: utf-8 -*-
"""
Transforme un cahier EST (Cahier_<niv>_ST_CORRIGE.docx) en LIVRET D'ACTIVITES.

Principe directeur (etabli sur Bord_Tle_A, 02/09/2026) :
    « ce qui NOMME reste, ce qui EXPLIQUE part »
L'apprenant garde le sujet de la regle et l'ecrit lui-meme ; il ne lit plus la
reponse toute faite. Les textes d'auteur ne sont JAMAIS touches (regle verbatim).

Usage :  python build_livret.py <2nde|1ere|Tle>
Sortie :  ../cahiers/Cahier_<niv>_ST_LIVRET_ACTIVITES.docx
"""
import sys, re, copy, os
from docx import Document

DOT = '\u2026'
LIGNE = DOT * 58          # ligne pleine a completer
TROU  = DOT * 25          # trou de verbe dans un objectif

# --- styles a ne JAMAIS modifier (textes d'auteur + apparat d'examen) -------
INTOUCHABLES = {'Texte', 'Corpus', 'Source', 'T_Disc', 'Texte_Titre',
                'Consigne', 'Astuce', 'Repere', 'Savais', 'Exercice',
                'T_Seq', 'T_Sem', 'T_Comp', 'T_Title'}

# --------------------------------------------------------------------------
def set_text(p, txt):
    """Reecrit un paragraphe en conservant le format du 1er run."""
    runs = p.runs
    if not runs:
        p.add_run(txt); return
    runs[0].text = txt
    for r in runs[1:]:
        r._element.getparent().remove(r._element)

def para_after(p, txt, style):
    """Insere un nouveau paragraphe juste apres p."""
    new = copy.deepcopy(p._element)
    p._element.addnext(new)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new, p._parent)
    for r in np.runs[1:]:
        r._element.getparent().remove(r._element)
    try:
        np.style = style
    except KeyError:
        pass
    if np.runs:
        np.runs[0].text = txt
    else:
        np.add_run(txt)
    return np

def n_lignes(supprime):
    """Nombre de lignes a completer, proportionnel au texte retire."""
    return max(1, min(4, round(len(supprime) / 90) + 1))

# --------------------------------------------------------------------------
# T1 — OBJECTIFS : le verbe d'action devient un trou
DETERM = ("le la les l' un une des du de d' son sa ses ce cet cette ces mon ma "
          "mes ton ta tes leur leurs a au aux en sur dans pour par avec entre "
          "chez vers selon comment pourquoi quand si que qu' plusieurs deux "
          "trois quatre cinq six").split()
INF = re.compile(r"^[a-zà-ÿ]+(er|ir|re|oir|oire)$", re.I)

def trouer_objectif(txt):
    m = re.search(r"(capable de\s+|capable d['\u2019]\s*)", txt, re.I)
    if not m:
        return None
    tete, reste = txt[:m.end()], txt[m.end():]
    toks = reste.split(' ')
    pris, i = [], 0
    while i < len(toks):
        mot = toks[i].strip(",;").lower().replace('\u2019', "'")
        if mot in DETERM or not mot:
            break
        if INF.match(mot) or mot in ('etre', '\u00eatre', 'faire', 'produire'):
            pris.append(toks[i]); i += 1
            # enchaine ", verbe" ou "et verbe"
            if i < len(toks) and toks[i].lower() in ('et', 'ou'):
                nxt = toks[i+1].strip(",;").lower() if i+1 < len(toks) else ''
                if INF.match(nxt):
                    pris.append(toks[i]); i += 1; continue
            if pris and pris[-1].endswith(','):
                continue
            break
        break
    if not pris:
        return None
    garde = ' '.join(toks[i:])
    sep = ', '.join([TROU] * len([x for x in pris if not x.lower() in ('et', 'ou')]))
    return (tete + sep + (' ' + garde if garde else '')).strip()


# --- passage a la premiere personne (le livret parle a l'apprenant) --------
PERS = [
    (r"l['’]apprenant sera capable d['’]", "je serai capable d'"),
    (r"l['’]apprenant sera capable de",         "je serai capable de"),
    (r"l['’]élève sera capable d['’]",    "je serai capable d'"),
    (r"l['’]élève sera capable de",            "je serai capable de"),
    (r"l['’]apprenant devra",                   "je devrai"),
    (r"l['’]élève devra",                      "je devrai"),
    (r"l['’]apprenant produira",                "je produirai"),
    (r"l['’]élève produira",                   "je produirai"),
    (r"À la fin de la leçon",                        "À la fin de cette leçon"),
    (r"de son cahier", "de mon cahier"),
]

def personnaliser(t):
    for pat, rep in PERS:
        t = re.sub(pat, rep, t, flags=re.I)
    return t

# --------------------------------------------------------------------------
# T2 — JE RETIENS : on garde le sujet + le verbe definitoire, on efface la suite
PIVOT = re.compile(
    r"\b(c['\u2019]est|ce sont|est|sont|d\u00e9signe|d\u00e9signent|consiste|consistent|"
    r"s['\u2019]appelle|se d\u00e9finit|sert \u00e0|servent \u00e0|permet de|permettent de|"
    r"n['\u2019]est pas|ne sont pas|signifie|correspond \u00e0|repose sur)\b", re.I)

def trouer_regle(txt):
    t = txt.strip()
    if len(t) < 60:
        return None                       # trop court : c'est un intitule
    m = PIVOT.search(t)
    if m and m.end() < len(t) * 0.65:
        garde = t[:m.end()].rstrip()
        return garde + ' ' + DOT * max(10, 58 - len(garde) % 58), t[m.end():]
    # pas de pivot : on garde la premiere proposition courte
    m2 = re.search(r"^(.{12,58}?)\s*[:,\u2013\u2014]\s", t)
    if m2:
        garde = m2.group(1).rstrip(' :,')
        return garde + ' : ' + DOT * 40, t[m2.end():]
    mots = t.split(' ')
    garde = ' '.join(mots[:5])
    return garde + ' ' + DOT * 40, ' '.join(mots[5:])

# --------------------------------------------------------------------------
# T3 — DEFINITIONS : le terme reste, la definition part
def trouer_definition(txt):
    t = txt.strip()
    m = re.search(r"^(.{0,80}?)\s*:\s", t)
    if not m or len(t) - m.end() < 40:
        return None
    return m.group(1).rstrip() + ' : ' + DOT * 34, t[m.end():]

# --------------------------------------------------------------------------
def transformer(src, dst):
    d = Document(src)
    st = {'obj': 0, 'regle': 0, 'defi': 0, 'lignes': 0, 'titres': 0}
    for p in list(d.paragraphs):
        s = p.style.name
        if s in INTOUCHABLES:
            continue
        t = p.text

        if s == 'Objectif':
            n = trouer_objectif(t)
            if n:
                set_text(p, personnaliser(n)); st['obj'] += 1

        elif s == 'Retient_Corps':
            r = trouer_regle(t)
            if r:
                garde, oté = r
                set_text(p, garde)
                anc = p
                for _ in range(n_lignes(oté)):
                    anc = para_after(anc, LIGNE, 'Retient_Corps'); st['lignes'] += 1
                st['regle'] += 1

        elif s == 'Definition':
            r = trouer_definition(t)
            if r:
                garde, oté = r
                set_text(p, garde)
                anc = p
                for _ in range(n_lignes(oté)):
                    anc = para_after(anc, LIGNE, 'Definition'); st['lignes'] += 1
                st['defi'] += 1

        elif s == 'Retient_Titre':
            if 'main' not in t:
                set_text(p, "\U0001f9e0 Je retiens \u2014 la r\u00e8gle s'\u00e9crit de ma main")
                st['titres'] += 1

        elif s == 'Reponse' and not t.strip():
            set_text(p, LIGNE); st['lignes'] += 1

    d.save(dst)
    return st

if __name__ == '__main__':
    niv = sys.argv[1] if len(sys.argv) > 1 else '2nde'
    base = os.path.join(os.path.dirname(__file__), '..', 'cahiers')
    src = os.path.join(base, 'Cahier_%s_ST_CORRIGE.docx' % niv)
    dst = os.path.join(base, 'Cahier_%s_ST_LIVRET_ACTIVITES.docx' % niv)
    st = transformer(src, dst)
    print('%s -> %s' % (os.path.basename(src), os.path.basename(dst)))
    print('   objectifs a trous : %(obj)d | regles ouvertes : %(regle)d | '
          'definitions ouvertes : %(defi)d | lignes a completer : %(lignes)d | '
          'titres : %(titres)d' % st)
