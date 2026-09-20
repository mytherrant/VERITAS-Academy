#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tools/cahiers_edition1.py — « MON CAHIER DE FRANÇAIS », 1ʳᵉ ÉDITION (2nde, 1ʳᵉ, Tˡᵉ)

    python tools/cahiers_edition1.py --charge "~/Desktop/veritas-ftp"
    python tools/cahiers_edition1.py --controle          # ne produit rien, mesure
    python tools/cahiers_edition1.py --seulement 1ere

CE QUE C'EST
  Trois ouvrages NEUFS, vendus 1 500 F chacun, À CÔTÉ des Bords déjà en ligne
  (`bord-2nde`, `bord-1ere`, `bord-tle`) — ce sont d'autres éditions, pas des
  mises à jour. Décision de Jacques du 19/09/2026 : « crée de nouveaux
  documents sans bouleverser ce qui existe déjà en ligne ». D'où des slugs à
  part, `bord-<classe>-ed1`, et un script à part : rien de ce qu'écrit
  `normaliser_cahiers.py` n'est touché, ni sa table, ni ses charges.

  Sources : `~/Downloads/1ère édition/Bord final <classe>.docx` + couverture.

POURQUOI UN LECTEUR DE .DOCX ET PAS LE CONVERTISSEUR DES BORDS
  Les Bords en ligne viennent de `content*.js`, déjà découpés en blocs par la
  maquette. Ces éditions-ci n'existent qu'en Word, et en Word MIS EN PAGE pour
  l'imprimeur, pas structuré : 95 % des paragraphes sont en style « Normal ».
  La structure est dans d'autres signes, qu'on lit un par un :
    · les titres de leçon et de domaine sont dans des ZONES DE TEXTE
      (`w:txbxContent`, en double : Choice + Fallback) — invisibles pour
      `paragraph.text` ;
    · « Leçon N : », « L'essentiel », « Je m'exerce », « Traitement : »,
      « Exercice N : », « Épreuve N », « Séquence N : » sont des libellés fixes ;
    · les exercices sont souvent posés dans des TABLEAUX DE MISE EN PAGE (une
      ligne, deux colonnes) : on les déplie, on ne les rend pas en tableau ;
      seuls les vrais tableaux (≥ 2 lignes et ≥ 2 colonnes) restent des tableaux ;
    · les pictogrammes décoratifs sont la MÊME image répétée des dizaines de
      fois ; une image qui n'apparaît qu'une fois est un document de l'exercice.

CE QUI DEVIENT UN CHAMP DE RÉPONSE
  Une question numérotée (« 1- », « 2. », « a) ») ou posée sous « Traitement »,
  et les lignes de pointillés du cahier. Le reste se lit. Semer des champs
  partout ferait un formulaire, pas un cahier.

OÙ VA QUOI
  --charge <dossier> → <dossier>/uploads/protected/livrets/booklet-<slug>.js   (FTP)
  livrets/extrait-<slug>.js                                                  (dépôt)
  uploads/oeuvres/livret_<slug>.jpg + _v.jpg                                  (dépôt)
  api/data/livrets_catalogue.json                                             (fiche)
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import re
import sys
from pathlib import Path

for _f in (sys.stdout, sys.stderr):          # la console Windows est en cp1252
    try:
        _f.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import docx                                   # python-docx
from docx.table import Table
from docx.text.paragraph import Paragraph

RACINE = Path(__file__).resolve().parent.parent
SOURCE = Path.home() / "Downloads" / "1ère édition"
CATALOGUE = RACINE / "api" / "data" / "livrets_catalogue.json"
EXTRAITS = RACINE / "livrets"
COUVERTURES = RACINE / "uploads" / "oeuvres"
PRIX = 1500                                   # « au prix de 1500 » — Jacques, 19/09/2026

OUVRAGES = {
    "bord-2nde-ed1": {"niveau": "2nde", "docx": "Bord final 2nde.docx",
                      "couv": "Couverture 2nde.jpg",
                      "titre": "Mon cahier de français 2ⁿᵈᵉ — 1ʳᵉ édition"},
    "bord-1ere-ed1": {"niveau": "1ere", "docx": "Bord final 1ère.docx",
                      "couv": "Couverture 1ère.jpg",
                      "titre": "Mon cahier de français 1ʳᵉ — 1ʳᵉ édition"},
    "bord-tle-ed1":  {"niveau": "tle", "docx": "Bord final TleA.docx",
                      "couv": "Couverture Tle.jpg",
                      "titre": "Mon cahier de français Tˡᵉ — 1ʳᵉ édition"},
}

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ── Petits outils ───────────────────────────────────────────────────────────
def net(t: str) -> str:
    return re.sub(r"[ \t ]+", " ", (t or "").replace("​", "")).strip()


def texte_zones(p) -> list[str]:
    """Le texte des zones de texte d'un paragraphe, dédoublonné (Word écrit la
    même zone deux fois : `mc:Choice` pour lui, `mc:Fallback` pour les autres)."""
    vus: list[str] = []
    for bx in p._p.iter(W + "txbxContent"):
        lignes = [net("".join(t.text or "" for t in q.iter(W + "t"))) for q in bx.iter(W + "p")]
        s = " ".join(x for x in lignes if x)
        if s and s not in vus:
            vus.append(s)
    return vus


def runs_de(p) -> list[dict]:
    """Les runs {t,b,i} du paragraphe, SANS le texte des zones de texte (déjà lu
    à part). Deux runs voisins de même habillage sont fusionnés."""
    out: list[dict] = []
    for r in p._p.iter(W + "r"):
        # un run qui vit dans une zone de texte n'appartient pas au paragraphe
        anc = r.getparent()
        dans_zone = False
        while anc is not None and anc is not p._p:
            if anc.tag == W + "txbxContent":
                dans_zone = True
                break
            anc = anc.getparent()
        if dans_zone:
            continue
        # ENFANTS DIRECTS seulement : le run qui porte une zone de texte la
        # contient (mc:AlternateContent → … → w:txbxContent), et `iter()` en
        # ramassait le texte — deux fois, Choice puis Fallback.
        t = "".join((x.text or "") if x.tag == W + "t" else ("\t" if x.tag == W + "tab" else "")
                    for x in r)
        if not t:
            continue
        rpr = r.find(W + "rPr")
        b = i = False
        if rpr is not None:
            bb = rpr.find(W + "b")
            ii = rpr.find(W + "i")
            b = bb is not None and bb.get(W + "val") not in ("0", "false")
            i = ii is not None and ii.get(W + "val") not in ("0", "false")
        run = {"t": t}
        if b:
            run["b"] = 1
        if i:
            run["i"] = 1
        if out and out[-1].get("b") == run.get("b") and out[-1].get("i") == run.get("i"):
            out[-1]["t"] += t
        else:
            out.append(run)
    # espaces : on normalise en gardant l'habillage
    for r in out:
        r["t"] = re.sub(r"[ \t ]+", " ", r["t"])
    while out and not out[0]["t"].strip():
        out.pop(0)
    while out and not out[-1]["t"].strip():
        out.pop()
    if out:
        out[0]["t"] = out[0]["t"].lstrip()
        out[-1]["t"] = out[-1]["t"].rstrip()
    return out


def texte_runs(runs: list[dict]) -> str:
    return net("".join(r.get("t", "") for r in runs))


def bloc_texte(y: str, runs: list[dict], **kw) -> dict:
    """`txt` quand le texte est nu, `r` quand il porte du gras ou de l'italique."""
    b = {"y": y}
    if any(r.get("b") or r.get("i") for r in runs):
        b["r"] = runs
    else:
        b["txt"] = texte_runs(runs)
    b.update(kw)
    return b


# ── Les libellés fixes du cahier ────────────────────────────────────────────
RE_PARTIE = re.compile(r"^(\d)\s*(?:ère|re|e|ème)\s+PARTIE\s*:\s*(.+)$", re.I)
RE_LECON = re.compile(r"^Le[çc]on\s*(\d+)\s*:?\s*(.*)$", re.I)
RE_ESSENTIEL = re.compile(r"^L[’']essentiel\b", re.I)
RE_EXERCE = re.compile(r"^Je m[’']exerce\b\s*:?\s*(.*)$", re.I)
RE_TRAITEMENT = re.compile(r"^Traitement\s*:?\s*(.*)$", re.I)
RE_EXERCICE = re.compile(r"^Exercice\s*(?::\s*)?(\d+)?\s*[:.]?\s*(.*)$", re.I)
RE_EPREUVE = re.compile(r"^(?:É|E)preuve\s*(\d+)?\s*:?\s*(.*)$", re.I)
RE_SEQUENCE = re.compile(r"^S[ée]quence\s*([\d/]+)\s*:\s*(.+)$", re.I)
RE_DOMAINE = re.compile(
    r"^(Communication|Morphosyntaxe|S[ée]mantique\s*/?\s*(?:Lexicologie)?|"
    r"Rh[ée]torique\s*/\s*Stylistique|Stylistique\s*/\s*Rh[ée]torique|"
    r"Lexicologie|Grammaire|Production [ée]crite|Essai|Contraction de texte)\s*[.:]?\s*$", re.I)
RE_GRAND = re.compile(
    r"^(EPREUVES? PAR S[ÉE]QUENCE|[ÉE]PREUVES? TYPES? EXAMEN|"
    r"[ÉE]PREUVES? DE PR[ÉE]PARATION AU BAC\w*)\s*[.:]?\s*$", re.I)
RE_LITT = re.compile(
    r"^(POUR ENTRER DANS L[’']ŒUVRE.*|VERS LE COMMENTAIRE COMPOS[ÉE]|"
    r"VERS LA DISSERTATION LITT[ÉE]RAIRE|PARTIE [A-H]\s*:.*)$", re.I)
RE_RUBRIQUE = re.compile(
    r"^(Axe\s*\d+\s*:?.*|Activit[ée]\s*\d+\s*:.*|[ÉE]valuation|[ÉE]l[ée]ments biographiques|"
    r"[ŒO]uvres? marquantes?|Autres textes sur la m[êe]me th[ée]matique.*|"
    r"Exercices? d[’']application|Comment s[’']y prendre\s*\??|Qu[’']est-ce que c[’']est\s*\??|"
    r"Texte\s*\d*\s*:.*|Consignes?\s*:?|Sujets?\s*\d*\s*:?|Application\s*:?|Remarques?\s*:?|"
    r"Approfondissement\s*:?)$", re.I)
RE_QUESTIONS = re.compile(r"^Questions?\s*:?\s*$", re.I)
RE_ZONE_OEUVRE = re.compile(r"^[ÉE]tude d[’']une? (?:de l[’'])?(?:[œo]uvre|œuvre).*", re.I)
RE_QUESTION = re.compile(r"^\s*(\d{1,2})\s*[-–.)°/]\s*(?=\S)(.*)$")
# « a) … », « a)Identifie … » (sans espace), « a- … » ; le numéro est en 1, le reste en 2.
RE_LETTRE = re.compile(r"^\s*([a-h])\s*(?:\)\s*|[\-–.]\s+)(?=\S)(.*)$")
RE_POINTILLES = re.compile(r"^[\s.…_·\-]{8,}$")
RE_FIN_POINTILLES = re.compile(r"\s*[.…_]{5,}\s*$")
VERBES = (r"Relève|Relevez|Cite|Citez|Explique|Expliquez|Justifie|Justifiez|Montre|Montrez|"
          r"Analyse|Analysez|Commente|Commentez|Rédige|Rédigez|Repère|Repérez|Compare|"
          r"Comparez|Résume|Résumez|Classe|Classez|Complète|Complétez|Indique|Indiquez|"
          r"Reformule|Reformulez|Recopie|Recopiez|Propose|Proposez|Imagine|Imaginez|"
          r"Identifie|Identifiez|Donne|Donnez|Dégage|Dégagez|Décode|Détermine|Déduis|"
          r"Illustre|Ressors|Construis|Souligne|Transforme|Réécris|Élabore|Elabore|Précise|"
          r"Dis|Trouve|Observe|Produis|Présente|Énumère|Enumère|Interprète|Mets|Étudie|Relis|"
          r"Établis|Etablis|Nomme|Situe|Formule|Examine|Définis|Explicite|Caractérise|"
          r"Conjugue|Remplace|Encadre|Associe|Relie|Choisis|Coche|Écris|Ecris|Transcris|"
          r"Distingue|Délimite|Relève|Recense|Fais|Emploie|Utilise|Construis|Explique")
RE_CONSIGNE = re.compile(r"^\s*(?:%s)\b" % VERBES)
RE_CONSIGNE_I = re.compile(r"^\s*(?:%s)\b" % VERBES, re.I)   # après un numéro : « 1. élabore »

# Les titres imprimés en capitales ont perdu leurs accents (« LES ELEMENTS DE
# PHONETIQUE »). Passés en minuscules, l'absence se VOIT : on les rend.
ACCENTS = {
    "elements": "éléments", "phonetique": "phonétique", "varietes": "variétés",
    "semantique": "sémantique", "rhetorique": "rhétorique", "specialise": "spécialisé",
    "neologie": "néologie", "tonalites": "tonalités", "litteraires": "littéraires",
    "litteraire": "littéraire", "litterature": "littérature", "theâtral": "théâtral",
    "theatral": "théâtral", "poetique": "poétique", "preparatoire": "préparatoire",
    "prèparatoire": "préparatoire", "redaction": "rédaction", "coherence": "cohérence",
    "cohesion": "cohésion", "epreuves": "épreuves", "epreuve": "épreuve",
    "sequence": "séquence", "methodologie": "méthodologie", "francaise": "française",
    "systeme": "système", "reduction": "réduction", "generalites": "généralités",
    "caracteristiques": "caractéristiques", "resume": "résumé", "etude": "étude",
    "oeuvre": "œuvre", "integrale": "intégrale", "enonciation": "énonciation",
    "negative": "négative", "modalites": "modalités", "strategies": "stratégies",
    "preposition": "préposition", "precis": "précis", "problematiques": "problématiques",
    "illustrees": "illustrées", "reponses": "réponses", "reponse": "réponse",
    "hierarchiques": "hiérarchiques", "procedes": "procédés", "denotation": "dénotation",
    "thematiques": "thématiques", "thematique": "thématique", "progressions": "progressions",
    "reecriture": "réécriture", "periode": "période", "specifique": "spécifique",
}


def casse(t: str) -> str:
    """« LES FOCALISATIONS » → « Les focalisations ». Un titre déjà en casse
    mixte n'est pas touché : il dit peut-être un nom propre."""
    lettres = [c for c in t if c.isalpha()]
    if not lettres or sum(c.isupper() for c in lettres) / len(lettres) < 0.6:
        return t
    s = t.lower()
    s = re.sub(r"[a-zàâäçéèêëîïôöùûüœ']+",
               lambda m: ACCENTS.get(m.group(0), m.group(0)), s)
    s = re.sub(r"\b(partie|axe|activité) ([a-h])\b",
               lambda m: m.group(1) + " " + m.group(2).upper(), s)
    return s[:1].upper() + s[1:]
RE_PRODUCTION = re.compile(r"^\s*(?:Rédige|Rédigez|Produis|Compose|Écris|Ecris|Fais le commentaire|"
                           r"Fais la dissertation|Élabore un plan|Elabore un plan)\b", re.I)


# ── Les images : un document, ou un pictogramme ─────────────────────────────
class Images:
    def __init__(self, doc):
        self.doc = doc
        self.compte: dict[str, int] = {}
        for rid in doc.element.body.xpath(".//a:blip/@r:embed"):
            h = self._hash(rid)
            if h:
                self.compte[h] = self.compte.get(h, 0) + 1
        self.vues: set[str] = set()

    def _blob(self, rid):
        part = self.doc.part.related_parts.get(rid)
        return getattr(part, "blob", None)

    def _hash(self, rid):
        b = self._blob(rid)
        return hashlib.md5(b).hexdigest() if b else None

    def documents(self, p) -> list[dict]:
        """Les images du paragraphe qui sont des DOCUMENTS : apparues une seule
        fois dans l'ouvrage. Un pictogramme (loupe, crayon, bandeau) revient
        des dizaines de fois — on le laisse au papier."""
        out = []
        for rid in p._p.xpath(".//a:blip/@r:embed"):
            h = self._hash(rid)
            if not h or self.compte.get(h, 0) != 1 or h in self.vues:
                continue
            self.vues.add(h)
            src = embarquer(self._blob(rid))
            if src:
                out.append({"y": "image", "src": src})
        return out


def embarquer(blob: bytes) -> str | None:
    """Image → JPEG 900 px maximum, en `data:` (le lecteur n'a pas de dossier
    d'images derrière le verrou ; embarquée, elle voyage avec le cahier)."""
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        im = Image.open(io.BytesIO(blob))
        im.load()
    except Exception:
        return None                                   # EMF/WMF : pas pour un navigateur
    if im.width < 80 or im.height < 80:
        return None                                   # une puce, pas un document
    if im.mode in ("RGBA", "LA", "P"):
        fond = Image.new("RGB", im.size, (255, 255, 255))
        im = im.convert("RGBA")
        fond.paste(im, mask=im.split()[-1])
        im = fond
    else:
        im = im.convert("RGB")
    im.thumbnail((900, 900))
    for q in (72, 60, 50):
        tampon = io.BytesIO()
        im.save(tampon, "JPEG", quality=q, optimize=True, progressive=True)
        if tampon.tell() <= 110_000:
            break
    return "data:image/jpeg;base64," + base64.b64encode(tampon.getvalue()).decode("ascii")


# ── Le lecteur ──────────────────────────────────────────────────────────────
class Lecteur:
    """Parcourt le corps du document DANS L'ORDRE (paragraphes et tableaux
    entremêlés) et émet des blocs `CAHIER_BLOCS`.

    L'état qui compte est `mode` :
      lire       texte courant ;
      texte      on est entre « Exercice N » et « Traitement » : un corpus ;
      questions  on est sous « Traitement » : chaque ligne est une question.
    """

    def __init__(self, doc):
        self.doc = doc
        self.images = Images(doc)
        self.blocs: list[dict] = []
        self.mode = "lire"
        self.titre_attente = ""          # le titre de leçon lu dans une zone de texte
        self.corpus: list[dict] = []     # runs du corpus en cours
        self.nlecons = 0
        self.nsections = 0
        self.partie = 0                  # 1 langue, 2 littérature, 3 méthodologie
        self.nq = 0                     # questions posées (pour le bilan)
        self.commence = False

    # ── émission ──
    def poser(self, b: dict):
        self.blocs.append(b)

    def fermer_corpus(self):
        if not self.corpus:
            return
        lignes = self.corpus
        self.corpus = []
        # La dernière ligne courte qui nomme un auteur et une date est la SOURCE.
        source = None
        dern = texte_runs(lignes[-1])
        if len(lignes) > 1 and len(dern) <= 160 and not RE_CONSIGNE_I.match(dern) \
                and not re.search(r"\b(dégage|relève|identifie|justifie)\b", dern, re.I) and (
                re.search(r"\b1[5-9]\d\d\b|\b20[0-2]\d\b", dern) or dern.count(",") >= 1
                and len(dern) < 110 and not dern.endswith((".", "?", "!", "»")) ):
            source = dern
            lignes = lignes[:-1]
        runs: list[dict] = []
        for k, l in enumerate(lignes):
            if k:
                runs.append({"br": 1})
            runs.extend(l)
        if runs:
            self.poser({"y": "corpus", "r": runs})
        if source:
            self.poser({"y": "source", "txt": source})

    def question(self, runs: list[dict], texte: str):
        texte = RE_FIN_POINTILLES.sub("", texte)
        m = RE_QUESTION.match(texte) or RE_LETTRE.match(texte)
        b = {"y": "question"}
        if m:
            b["no"] = m.group(1)
            corps = net(m.group(2))
        else:
            corps = texte
        if not corps:
            return
        b["txt"] = corps
        self.poser(b)
        self.nq += 1
        if RE_PRODUCTION.match(corps):
            self.poser({"y": "lines", "n": 10})

    def est_question(self, t: str) -> bool:
        """Une ligne POSE-t-elle une question ? Numérotée avec un verbe de
        consigne ou un « ? », ou lettrée (« a) »), ou verbe de consigne en tête."""
        m = RE_QUESTION.match(t)
        if m and (RE_CONSIGNE_I.match(m.group(2)) or t.endswith("?")):
            return True
        # Hors méthodologie, un item numéroté est une consigne (« 1- Situe le
        # texte », « 2- Formule 4 hypothèses ») ; en méthodologie, la liste
        # numérotée EXPLIQUE (« 2- Lorsque le domaine d'application est
        # fermé… ») et reste à lire.
        if m and self.partie < 3 and len(t) < 400:
            return True
        if RE_LETTRE.match(t) and len(t) < 400:
            return True
        return bool(RE_CONSIGNE.match(t)) and len(t) < 300 and not t.endswith(":")

    def invite(self) -> bool:
        """Le bloc précédent appelle-t-il des réponses ? (« réponds aux
        questions suivantes : », une question déjà posée, une rubrique.)"""
        if not self.blocs:
            return False
        d = self.blocs[-1]
        y = d.get("y")
        if y in ("question", "lines", "rubrique", "rubriqueH", "source", "corpus", "table"):
            return True
        if y == "texte":
            s = d.get("txt") or texte_runs(d.get("r", []))
            return s.endswith(":") or bool(re.search(r"questions suivantes|r[ée]ponds", s, re.I))
        return False

    def paragraphe_texte(self, t: str):
        """Du texte nu, rangé selon le mode courant (corpus ou lecture)."""
        if self.mode == "texte":
            self.corpus.append([{"t": t}])
        else:
            self.corpus.append([{"t": t}])      # la ligne d'un texte d'étude
            self.fermer_corpus()

    # ── un paragraphe ──
    def paragraphe(self, p):
        zones = texte_zones(p)
        runs = runs_de(p)
        t = texte_runs(runs)

        if not self.commence:
            if RE_PARTIE.match(t):
                self.commence = True
            else:
                return

        for z in zones:
            self.zone(z)

        imgs = self.images.documents(p)

        if not t or not re.search(r"\w", t):
            for im in imgs:
                self.fermer_corpus()
                self.poser(im)
            return

        # « … vous saurez quel remède apporter à mes maux. Traitement : » — le
        # libellé collé en fin de paragraphe : on coupe, puis on le traite.
        m = re.match(r"^(.{20,}?)\s*\bTraitement\s*:\s*$", t)
        if m:
            self.paragraphe_texte(m.group(1))
            self.fermer_corpus()
            self.mode = "questions"
            self.poser({"y": "rubriqueH", "txt": "Traitement"})
            return
        # « Texte 1 : Le vandalisme scolaire est… » : l'annonce ET le texte.
        m = re.match(r"^(Texte\s*\d*)\s*:\s*(.{60,})$", t)
        if m and self.partie < 3:
            self.fermer_corpus()
            self.mode = "texte"
            self.poser({"y": "rubrique", "txt": net(m.group(1))})
            self.corpus.append([{"t": net(m.group(2))}])
            return

        # ── les divisions ──
        m = RE_PARTIE.match(t)
        if m:
            self.fermer_corpus()
            self.mode = "lire"
            self.nsections += 1
            self.partie = int(m.group(1))
            self.poser({"y": "part", "no": m.group(1), "title": casse(net(m.group(2)))})
            return
        if RE_GRAND.match(t):
            self.fermer_corpus()
            self.mode = "lire"
            self.poser({"y": "section", "title": casse(t.rstrip(" .:"))})
            return
        m = RE_SEQUENCE.match(t)
        if m and len(t) < 260:
            self.fermer_corpus()
            self.mode = "lire"
            self.poser({"y": "lecon", "no": "S" + m.group(1),
                        "title": "Séquence " + m.group(1) + " — " + net(m.group(2))})
            self.nlecons += 1
            return
        m = RE_LECON.match(t)
        if m and len(t) < 160:
            self.fermer_corpus()
            self.mode = "lire"
            titre = casse(net(m.group(2))) or self.titre_attente or ("Leçon " + m.group(1))
            self.titre_attente = ""
            self.poser({"y": "lecon", "no": m.group(1), "title": titre})
            self.nlecons += 1
            return
        if RE_LITT.match(t):
            self.fermer_corpus()
            self.mode = "lire"
            self.poser({"y": "lecon", "title": casse(t)})
            self.nlecons += 1
            return
        m = RE_EPREUVE.match(t)
        if m and (m.group(1) or not m.group(2)) and len(t) < 300:
            self.fermer_corpus()
            self.mode = "texte"
            reste = net(m.group(2))
            self.poser({"y": "rubrique", "txt": ("Épreuve " + (m.group(1) or "")).strip()})
            if reste:
                self.corpus.append([{"t": reste, "i": 1}] if reste.startswith("(") else [{"t": reste}])
            return

        # ── les rubriques de la leçon ──
        if RE_ESSENTIEL.match(t) and len(t) < 40:
            self.fermer_corpus()
            self.mode = "lire"
            self.poser({"y": "rubrique", "txt": "L’essentiel"})
            return
        m = RE_EXERCE.match(t)
        if m and len(t) < 60:
            self.fermer_corpus()
            self.mode = "lire"
            self.poser({"y": "rubrique", "txt": "Je m’exerce"})
            return
        m = RE_TRAITEMENT.match(t)
        if m and len(t) < 400:
            self.fermer_corpus()
            self.mode = "questions"
            self.poser({"y": "rubriqueH", "txt": "Traitement"})
            if net(m.group(1)):
                self.question(runs, net(m.group(1)))
            return
        if RE_QUESTIONS.match(t):
            self.fermer_corpus()
            self.mode = "questions"
            self.poser({"y": "rubriqueH", "txt": "Questions"})
            return
        m = RE_EXERCICE.match(t)
        if m and (m.group(1) or len(t) < 30) and not RE_CONSIGNE.match(net(m.group(2)) or "x"):
            self.fermer_corpus()
            self.mode = "texte"
            titre = "Exercice " + (m.group(1) or "")
            reste = net(m.group(2))
            # « Exercice 1 : l'émetteur et le récepteur » : un sous-titre, pas
            # la première ligne du texte à étudier.
            if reste and len(reste) < 80 and not reste.endswith((".", "!", "?", "»")):
                titre += " — " + reste
                reste = ""
            self.poser({"y": "rubriqueH", "txt": titre.strip()})
            if reste:
                self.corpus.append([{"t": reste}])
            return
        if RE_DOMAINE.match(t):
            self.fermer_corpus()
            der = self.blocs[-1] if self.blocs else {}
            if der.get("y") == "section" and net(der.get("title", "")).lower() == t.rstrip(" .:").lower():
                return                          # le titre de la zone vient de le dire
            self.mode = "questions"
            self.poser({"y": "rubriqueH", "txt": t.rstrip(" .:")})
            return
        if RE_RUBRIQUE.match(t) and len(t) < 200 and (runs and runs[0].get("b")):
            self.fermer_corpus()
            # « Texte 1 : » annonce un corpus ; les autres rubriques ramènent au texte courant
            self.mode = "texte" if re.match(r"^Texte\b", t, re.I) else "lire"
            self.poser({"y": "rubrique", "txt": casse(t.rstrip(" :"))})
            return

        # ── les lignes à remplir ──
        if RE_POINTILLES.match(t):
            self.fermer_corpus()
            if self.blocs and self.blocs[-1].get("y") == "lines":
                self.blocs[-1]["n"] = min(20, self.blocs[-1]["n"] + 1)
            elif self.blocs and self.blocs[-1].get("y") == "question":
                pass                            # la question porte déjà sa boîte
            else:
                self.poser({"y": "lines", "n": 1})
            return

        # ── le contenu ──
        if self.mode == "questions":
            # Sous « Traitement » : ce qui pose une question en est une ; la
            # PREMIÈRE ligne l'est toujours (« Même consigne que l'exercice
            # précédent ») ; le reste est une explication à lire.
            premiere = self.blocs and self.blocs[-1].get("y") == "rubriqueH"
            if self.est_question(t) or t.endswith("?") or (premiere and len(t) < 260):
                self.question(runs, t)
            else:
                self.poser(bloc_texte("texte", runs))
        elif self.mode == "texte":
            # une question égarée au milieu d'un corpus reste une question
            m = RE_QUESTION.match(t)
            if m and (RE_CONSIGNE_I.match(m.group(2)) or t.endswith("?")):
                self.fermer_corpus()
                self.mode = "questions"
                self.question(runs, t)
            else:
                self.corpus.append(runs)
                # La référence (« Marcel Frydman, Jeunes, éducation et violence,
                # 2008. ») FERME le texte : ce qui suit est une consigne, même
                # sans « Traitement » pour l'annoncer.
                if len(self.corpus) > 1 and len(t) <= 160 and not t.endswith("?") \
                        and re.search(r"\b1[5-9]\d\d\b|\b20[0-2]\d\b", t) and "," in t:
                    self.fermer_corpus()
                    self.mode = "lire"
        else:
            # « NAM : Que veux-tu dire ? » est une réplique, pas une consigne.
            replique = re.match(r"^[A-ZÉÈÀÂÎÔÛÇ][A-ZÉÈÀÂÎÔÛÇ' -]{1,30}\s*:", t)
            if self.est_question(t) or (t.endswith("?") and len(t) < 260 and not replique
                                        and self.partie < 3 and self.invite()):
                self.question(runs, t)
            else:
                self.poser(bloc_texte("texte", runs))
        for im in imgs:
            self.fermer_corpus()
            self.poser(im)

    def zone(self, z: str):
        """Une zone de texte : titre de leçon, de domaine, d'œuvre — ou encadré."""
        z = net(z)
        if not z:
            return
        if len(z) <= 90:
            self.fermer_corpus()
            self.mode = "lire"
            if RE_ZONE_OEUVRE.match(z):
                self.poser({"y": "section", "title": re.sub(r"\s*:\s*", " : ", z, count=1)})
                return
            if re.match(r"^Texte\s*\d+\b", z, re.I) and self.partie < 3:
                # un texte d'étude annoncé en zone : les lignes suivantes sont son corpus
                self.poser({"y": "rubrique", "txt": z})
                self.mode = "texte"
                return
            m = re.match(r"^([IVX]+)\s*[-–.]\s*(.+)$", z)
            if m:
                self.poser({"y": "section", "no": m.group(1), "title": casse(m.group(2))})
                return
            if z.isupper() and len(z.split()) <= 3 and self.blocs \
                    and self.blocs[-1].get("y") == "part":
                self.poser({"y": "section", "title": casse(z)})
                return
            # un titre de leçon : il précède « Leçon N : »
            self.titre_attente = casse(z)
            if not RE_LITT.match(z):
                # s'il n'est pas suivi d'un « Leçon N », il reste un intertitre
                self.poser({"y": "_titre", "txt": self.titre_attente})
            return
        self.fermer_corpus()
        self.poser({"y": "boxHead", "txt": "À retenir"})
        self.poser({"y": "boxBody", "txt": z})

    # ── un tableau ──
    def tableau(self, tb: Table):
        lignes = []
        for r in tb.rows:
            vus, cells = set(), []
            for c in r.cells:
                if id(c._tc) in vus:
                    continue
                vus.add(id(c._tc))
                cells.append(c)
            lignes.append(cells)
        nl = len(lignes)
        nc = max((len(x) for x in lignes), default=0)
        if nl >= 2 and nc >= 2:
            self.fermer_corpus()
            rows = []
            for cells in lignes:
                rows.append([{"t": net(" ".join(p.text for p in c.paragraphs))} for c in cells])
            self.poser({"y": "table", "rows": rows})
            for cells in lignes:
                for c in cells:
                    for p in c.paragraphs:
                        for im in self.images.documents(p):
                            self.poser(im)
            return
        # Tableau de MISE EN PAGE : on le déplie, cellule après cellule.
        for cells in lignes:
            for c in cells:
                for el in c._tc.iterchildren():
                    tag = el.tag.split("}")[1]
                    if tag == "p":
                        self.paragraphe(Paragraph(el, c))
                    elif tag == "tbl":
                        self.tableau(Table(el, c))
                self.fermer_corpus()

    def lire(self):
        for el in self.doc.element.body.iterchildren():
            tag = el.tag.split("}")[1]
            if tag == "p":
                self.paragraphe(Paragraph(el, self.doc))
            elif tag == "tbl" and self.commence:
                self.tableau(Table(el, self.doc))
        self.fermer_corpus()
        return finir(self.blocs)


def finir(blocs: list[dict]) -> list[dict]:
    """Dernière passe : un intertitre de zone suivi de « Leçon N » a déjà servi
    de titre à la leçon — on le retire ; sinon il devient un intertitre."""
    out: list[dict] = []
    for k, b in enumerate(blocs):
        if b.get("y") == "_titre":
            suiv = blocs[k + 1] if k + 1 < len(blocs) else {}
            if suiv.get("y") == "lecon" and suiv.get("title") == b["txt"]:
                continue
            out.append({"y": "rubrique", "txt": b["txt"]})
            continue
        out.append(b)
    # Rubriques en double d'affilée (« Je m'exerce » puis « Je m'exerce »)
    propre: list[dict] = []
    for b in out:
        if propre and b.get("y") in ("rubrique", "rubriqueH") and propre[-1] == b:
            continue
        propre.append(b)
    return propre


# ── Extrait gratuit : deux leçons éloignées ─────────────────────────────────
def extraire(blocs: list[dict]) -> list[dict]:
    """Même règle que les autres cahiers (Jacques, 27/08) : deux leçons au plus,
    non suivies, sans image embarquée (la CI plafonne un extrait à 120 Ko)."""
    lecons, entete, cour = [], [], None
    for b in blocs:
        y = b.get("y")
        if y in ("part", "section"):
            entete = [x for x in entete if x.get("y") != y] + [b]
            continue
        if y in ("lecon", "epreuve"):
            cour = {"entete": list(entete), "blocs": [b]}
            lecons.append(cour)
            continue
        if cour is not None:
            cour["blocs"].append(b)
    def poids(l):
        return sum(1 for b in l["blocs"] if b.get("y") in ("question", "lines"))
    utiles = [l for l in lecons if poids(l) >= 3 and len(l["blocs"]) >= 8]
    if len(utiles) < 2:
        utiles = lecons
    a = utiles[len(utiles) // 4]
    b2 = utiles[(len(utiles) * 2) // 3]
    if utiles.index(b2) - utiles.index(a) < 2:
        b2 = utiles[-1]
    out = []
    for n, l in enumerate([a, b2]):
        if n:
            out.append({"y": "rubrique", "txt": "⋯ plus loin dans le cahier ⋯"})
        out.extend(l["entete"])
        for b in l["blocs"]:
            if b.get("y") == "image":
                continue
            out.append(b)
    return out


def js(blocs: list[dict]) -> str:
    return ("window.CAHIER_BLOCS=" +
            json.dumps(blocs, ensure_ascii=False, separators=(",", ":")) + ";\n")


def couverture(src: Path, slug: str) -> None:
    """Grande (1200×1600) et vignette (460×613), comme les autres ouvrages.

    ⚠️ ON PART DE LA COUVERTURE NETTOYÉE quand elle existe. Les fichiers reçus
    portent le cachet « 3000 FCFA » — le prix du cahier de PAPIER — et, pour la
    2ⁿᵈᵉ, un bouton d'application resté dans une capture d'écran. Afficher
    3 000 F sur l'image d'un produit vendu 1 500 F, c'est demander à l'acheteur
    lequel des deux prix est le vrai. `tools/couvertures_edition1.py` produit
    les versions nettoyées ; sans elles, on prend l'original plutôt que de ne
    rien produire, mais on le DIT.
    """
    from PIL import Image
    propre = RACINE / "uploads" / "oeuvres" / "_sources" / f"couverture-{slug.replace('bord-', '').replace('-ed1', '')}-ed1.jpg"
    if propre.is_file():
        src = propre
    else:
        print(f"    ⚠️  {slug} : couverture NON nettoyée (cachet « 3000 FCFA » encore "
              f"visible) — lancer d'abord tools/couvertures_edition1.py")
    im = Image.open(src).convert("RGB")
    for suffixe, taille in (("", (1200, 1600)), ("_v", (460, 613))):
        c = im.copy()
        c.thumbnail(taille, Image.LANCZOS)
        fond = Image.new("RGB", taille, (255, 255, 255))
        fond.paste(c, ((taille[0] - c.width) // 2, (taille[1] - c.height) // 2))
        fond.save(COUVERTURES / f"livret_{slug}{suffixe}.jpg", "JPEG", quality=86,
                  optimize=True, progressive=True)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--charge", type=Path, help="dossier de la charge FTP (~/Desktop/veritas-ftp)")
    ap.add_argument("--source", type=Path, default=SOURCE)
    ap.add_argument("--seulement", action="append", default=[])
    ap.add_argument("--controle", action="store_true", help="ne rien écrire")
    a = ap.parse_args()

    slugs = [s for s in OUVRAGES if not a.seulement
             or s in a.seulement or s.split("-")[1] in a.seulement]

    sortie = None
    if a.charge and not a.controle:
        sortie = Path(str(a.charge).replace("~", str(Path.home()), 1)) / "uploads" / "protected" / "livrets"
        # Le bon dépôt FTP porte déjà des dizaines de cahiers (voir la mémoire
        # « deux dossiers veritas-ftp ») : on le constate avant d'y écrire.
        deja = list(sortie.glob("booklet-*.js")) if sortie.is_dir() else []
        if len(deja) < 5:
            print(f"✗ {sortie} ne contient que {len(deja)} cahier(s) : ce n'est "
                  f"probablement pas le dépôt FTP. Rien n'est écrit.")
            return 1

    cat = json.loads(CATALOGUE.read_text(encoding="utf-8"))
    faits = 0
    for slug in slugs:
        spec = OUVRAGES[slug]
        f = a.source / spec["docx"]
        if not f.is_file():
            print(f"  ✗ {slug} : source absente ({f})")
            continue
        lec = Lecteur(docx.Document(str(f)))
        blocs = lec.lire()
        tete = [{"y": "rubrique", "txt": spec["titre"]}]
        blocs = tete + blocs
        ext = extraire(blocs)
        contenu, apercu = js(blocs), js(ext)
        types: dict[str, int] = {}
        for b in blocs:
            types[b["y"]] = types.get(b["y"], 0) + 1
        print(f"  {slug:14} {len(blocs):5} blocs · {lec.nq:4} questions · "
              f"{types.get('lecon', 0):3} leçons · {types.get('image', 0):2} images · "
              f"{len(contenu) / 1024:6.0f} Ko · aperçu {len(apercu) / 1024:4.0f} Ko")
        print("                 " + json.dumps(types, ensure_ascii=False))
        if a.controle:
            continue
        if sortie:
            (sortie / f"booklet-{slug}.js").write_text(contenu, encoding="utf-8")
        (EXTRAITS / f"extrait-{slug}.js").write_text(apercu, encoding="utf-8")
        couverture(a.source / spec["couv"], slug)
        cat["ouvrages"][slug] = {
            "titre": spec["titre"], "niveau": spec["niveau"], "mode": "interactif",
            # Pas de guide pour ces éditions : un code « guide » n'ouvrirait rien.
            "kinds": ["livret"], "prix": PRIX, "prixGuide": 0,
            "pages": 0, "pagesLibres": 0,
        }
        faits += 1

    if faits:
        CATALOGUE.write_text(json.dumps(cat, ensure_ascii=False, indent=2) + "\n",
                             encoding="utf-8")
        print(f"\n  catalogue : {len(cat['ouvrages'])} ouvrages")
        if not sortie:
            print("  ⚠️  --charge absent : les cahiers ÉLÈVE n'ont pas été écrits.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
