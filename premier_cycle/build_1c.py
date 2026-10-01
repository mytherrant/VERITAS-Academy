# -*- coding: utf-8 -*-
"""Construit les quatre manuels d'étude d'œuvres du premier cycle.

    python premier_cycle/build_1c.py          # les quatre
    python premier_cycle/build_1c.py 6e       # un seul

Chaque manuel réunit **trois œuvres d'une même classe**, une section de
passerelles qui les fait se répondre, les corrigés de tous les jeux et de
toutes les épreuves, et la note aux enseignants. Le document est reconstruit
entièrement à chaque exécution : le contenu vit dans les modules Python, jamais
dans le `.docx`.
"""
import os
import sys

ICI = os.path.dirname(os.path.abspath(__file__))
if ICI not in sys.path:
    sys.path.insert(0, ICI)

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import kit
from kit import docxkit

SORTIE = r"C:\Users\Mythe Errant\Desktop\Manuels\FINAUX_1er_cycle"

MANUELS = {
    "6e": {"module": "contenu_6e", "classe": "Sixième",
           "titre": "Trois œuvres au programme de la classe de 6ᵉ"},
    "5e": {"module": "contenu_5e", "classe": "Cinquième",
           "titre": "Trois œuvres au programme de la classe de 5ᵉ"},
    "4e": {"module": "contenu_4e", "classe": "Quatrième",
           "titre": "Trois œuvres au programme de la classe de 4ᵉ"},
    "3e": {"module": "contenu_3e", "classe": "Troisième",
           "titre": "Trois œuvres au programme de la classe de 3ᵉ"},
}


def _page(doc):
    """Format A4, marges d'un cahier d'activités : large à gauche pour la
    reliure, assez d'air à droite pour que l'élève annote."""
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin = Cm(2.3), Cm(1.8)
    s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)


def _couverture(doc, classe, titre, oeuvres):
    for texte, taille, gras, couleur, avant in [
            ("CENTRE VÉRITAS", 13, True, docxkit.OR, 60),
            ("Étude d'œuvres intégrales", 26, True, docxkit.BLEU, 18),
            ("Classe de " + classe, 20, False, docxkit.BLEU_CLAIR, 6),
            ("", 11, False, docxkit.GRIS, 20)]:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(avant)
        r = par.add_run(texte)
        r.font.size = Pt(taille)
        r.font.bold = gras
        r.font.color.rgb = couleur
    for i, (t, a) in enumerate(oeuvres, 1):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(10)
        r = par.add_run("%d.  %s" % (i, t))
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = docxkit.BLEU
        r2 = par.add_run("\n" + a)
        r2.font.size = Pt(11.5)
        r2.font.italic = True
        r2.font.color.rgb = docxkit.GRIS
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(50)
    r = par.add_run("© Mythe Errant · Centre VÉRITAS · veritas-school.com")
    r.font.size = Pt(9)
    r.font.color.rgb = docxkit.GRIS
    doc.add_page_break()


def construire(cle):
    conf = MANUELS[cle]
    mod = __import__(conf["module"])
    doc = docx.Document()
    _page(doc)
    _couverture(doc, conf["classe"], conf["titre"], mod.OEUVRES)

    blocs, corriges = mod.manuel()
    docxkit.render(doc, blocs)
    docxkit.render(doc, corriges)

    os.makedirs(SORTIE, exist_ok=True)
    chemin = os.path.join(SORTIE, "%s_Etude_oeuvres_integrales.docx" % cle)
    doc.save(chemin)
    return chemin, docxkit.words(blocs) + docxkit.words(corriges)


if __name__ == "__main__":
    cibles = sys.argv[1:] or list(MANUELS)
    for c in cibles:
        chemin, n = construire(c)
        print("%-3s  %6d mots  ->  %s" % (c, n, chemin))
